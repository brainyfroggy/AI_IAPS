from __future__ import annotations

import importlib.util
import re
import shutil
import sys
import urllib.request
from pathlib import Path


LOCAL_PYTHON_DEPS = Path(__file__).resolve().parent / ".python_deps"
if LOCAL_PYTHON_DEPS.exists():
    sys.path.insert(0, str(LOCAL_PYTHON_DEPS))

import nibabel as nib
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill


PROJECT_ROOT = Path(r"N:\Experimental_Data\yujunchen\projects\AI_IAPS")
ROI_TABLE_DIR = PROJECT_ROOT / "ROI table"
BASE_SCRIPT = ROI_TABLE_DIR / "compute_roi_overlap_tables.py"
CONTRAST_DIR = PROJECT_ROOT / "GLM" / "outputs_seclvl" / "28subs"
OUT_DIR = ROI_TABLE_DIR / "28subs_fdr_positive_single_contrast_rois_merged_mricrogl"
ATLAS_DIR = OUT_DIR / "atlas"
HCP_LONG_NAME_URL = (
    "https://bitbucket.org/dpat/tools/raw/master/REF/ATLASES/"
    "HCP-MMP1_UniqueRegionList.csv"
)
HCP_LONG_NAME_CSV = ATLAS_DIR / "HCP-MMP1_UniqueRegionList.csv"
MICROGL_SYSTEM_ATLAS_DIR = Path(r"C:\MRIcroGL\Resources\atlas")
MICROGL_AAL3_ATLAS = MICROGL_SYSTEM_ATLAS_DIR / "AAL3v1.nii.gz"
MICROGL_AAL3_LABELS = MICROGL_SYSTEM_ATLAS_DIR / "AAL3v1.nii.txt"
MICROGL_KASTNER_ATLAS = MICROGL_SYSTEM_ATLAS_DIR / "kastner.nii.gz"
MICROGL_KASTNER_LABELS = MICROGL_SYSTEM_ATLAS_DIR / "kastner.nii.txt"

CONTRASTS = {
    "Pleasant_vs_Neutral": CONTRAST_DIR / "fdr_tstat_cond_Pl-Nt.nii",
    "AI_Pleasant_vs_AI_Neutral": CONTRAST_DIR / "fdr_tstat_cond_AI_Pl-AI_Nt.nii",
    "Unpleasant_vs_Neutral": CONTRAST_DIR / "fdr_tstat_cond_Up-Nt.nii",
    "AI_Unpleasant_vs_AI_Neutral": CONTRAST_DIR / "fdr_tstat_cond_AI_Up-AI_Nt.nii",
}


def load_base_module():
    spec = importlib.util.spec_from_file_location("roi_overlap_base", BASE_SCRIPT)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not load {BASE_SCRIPT}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def merged_roi_name(roi_name: str) -> str:
    return re.sub(r"_\d+$", "", roi_name)


def safe_filename(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", name)
    cleaned = cleaned.strip("._")
    return cleaned or "ROI"


def atlas_short_name(atlas_name: str) -> str:
    if atlas_name.startswith("Schaefer"):
        return "Schaefer400"
    if atlas_name.startswith("HCP"):
        return "HCP_MMP1"
    if atlas_name.startswith("AAL3"):
        return "AAL3"
    if atlas_name.startswith("Kastner") or atlas_name.startswith("Wang"):
        return "KastnerWang"
    if atlas_name.startswith("HarvardOxford"):
        return "HarvardOxford"
    return safe_filename(atlas_name)


SCHAEFER_NETWORK_FULL_NAMES = {
    "Vis": "Visual Network",
    "SomMot": "Somatomotor Network",
    "DorsAttn": "Dorsal Attention Network",
    "SalVentAttn": "Salience Ventral Attention Network",
    "Limbic": "Limbic Network",
    "Cont": "Frontoparietal Control Network",
    "Default": "Default Mode Network",
}


def clean_readable_name(value: object) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value).replace("_", " ")).strip()


def load_hcp_region_info() -> dict[str, dict[str, str]]:
    if not HCP_LONG_NAME_CSV.exists():
        HCP_LONG_NAME_CSV.parent.mkdir(parents=True, exist_ok=True)
        urllib.request.urlretrieve(HCP_LONG_NAME_URL, HCP_LONG_NAME_CSV)

    df = pd.read_csv(HCP_LONG_NAME_CSV)
    info: dict[str, dict[str, str]] = {}
    for _, row in df.iterrows():
        hemi = str(row.get("LR", "")).strip()
        region = str(row.get("region", "")).strip()
        if not hemi or not region:
            continue
        label_name = f"{hemi}_{region}"
        official_name = clean_readable_name(row.get("regionLongName", label_name))
        cortex = clean_readable_name(row.get("cortex", ""))
        lobe = clean_readable_name(row.get("Lobe", ""))
        info[label_name] = {
            "official_name": official_name,
            "full_name": hcp_interpretable_name(official_name, region, hemi, cortex),
            "lobe": lobe,
            "cortex": cortex,
        }
    return info


def hcp_interpretable_name(
    official_name: str,
    region: str,
    hemi: str,
    cortex: str,
) -> str:
    side = {"L": "L", "R": "R"}.get(hemi, hemi)
    cortex_name = clean_readable_name(cortex)

    name_without_side = re.sub(rf"\s+{re.escape(side)}$", "", official_name).strip()
    name_without_area = re.sub(r"^Area\s+", "", name_without_side).strip()

    if not official_name.startswith("Area "):
        return official_name

    # HCP MMP names often use "Area <code>" as the official label. For ROI
    # tables, prepend the HCP cortex grouping so the label is interpretable.
    if cortex_name and name_without_area == region:
        return f"{cortex_name} - {region} {side}"
    if cortex_name and re.fullmatch(r"[A-Za-z0-9+\-.]+", name_without_area):
        return f"{cortex_name} - {name_without_area} {side}"
    return f"{name_without_area} {side}".strip()


def readable_merged_roi_full_name(
    merged_name: str,
    atlas_name: str,
    hcp_region_info: dict[str, dict[str, str]] | None = None,
) -> str:
    if atlas_name.startswith("HCP"):
        if hcp_region_info and merged_name in hcp_region_info:
            return hcp_region_info[merged_name]["full_name"]
        return merged_name

    if not atlas_name.startswith("Schaefer"):
        return clean_readable_name(merged_name)

    parts = merged_name.split("_")
    if len(parts) < 3:
        return merged_name

    hemisphere = {"LH": "Left", "RH": "Right"}.get(parts[1], parts[1])
    network = SCHAEFER_NETWORK_FULL_NAMES.get(parts[2], parts[2])
    remainder = "_".join(parts[3:])

    readable_parts = [hemisphere, network]
    if remainder:
        readable_parts.append(clean_readable_name(remainder))
    return " ".join(readable_parts)


def parse_mricrogl_labels(label_path: Path) -> dict[int, str]:
    labels: dict[int, str] = {}
    with label_path.open("r", encoding="utf-8-sig") as f:
        for line in f:
            stripped = line.strip()
            if not stripped:
                continue
            parts = stripped.split()
            if len(parts) < 2:
                continue
            try:
                roi_id = int(float(parts[0]))
            except ValueError:
                continue
            labels[roi_id] = parts[1]
    return labels


def copy_atlas_file(src: Path, dst: Path) -> None:
    if not src.exists():
        raise FileNotFoundError(f"Missing atlas file: {src}")
    dst.parent.mkdir(parents=True, exist_ok=True)
    if not dst.exists() or src.stat().st_mtime > dst.stat().st_mtime:
        shutil.copy2(src, dst)


def load_local_label_atlas(
    base,
    ref_img,
    atlas_name: str,
    source_atlas: Path,
    source_labels: Path,
    atlas_subdir: str,
    atlas_filename: str,
) -> tuple[str, nib.Nifti1Image, dict[int, str]]:
    atlas_path = ATLAS_DIR / atlas_subdir / atlas_filename
    if atlas_filename.endswith(".nii.gz"):
        label_filename = atlas_filename.removesuffix(".nii.gz") + ".nii.txt"
    else:
        label_filename = Path(atlas_filename).stem + ".nii.txt"
    labels_path = ATLAS_DIR / atlas_subdir / label_filename
    copy_atlas_file(source_atlas, atlas_path)
    copy_atlas_file(source_labels, labels_path)

    atlas_img = nib.load(str(atlas_path))
    labels = parse_mricrogl_labels(labels_path)
    resampled = base.image.resample_to_img(
        atlas_img,
        ref_img,
        interpolation="nearest",
        force_resample=True,
        copy_header=True,
    )
    return atlas_name, resampled, labels


def load_aal3(base, ref_img):
    return load_local_label_atlas(
        base,
        ref_img,
        "AAL3v1_MRIcroGL",
        MICROGL_AAL3_ATLAS,
        MICROGL_AAL3_LABELS,
        "AAL3",
        "AAL3v1.nii.gz",
    )


def load_kastner_wang(base, ref_img):
    return load_local_label_atlas(
        base,
        ref_img,
        "Kastner_Wang_25VisualROIs",
        MICROGL_KASTNER_ATLAS,
        MICROGL_KASTNER_LABELS,
        "Kastner_Wang",
        "kastner.nii.gz",
    )


def mricrogl_label_lines(
    labels: dict[int, str],
    hcp_region_info: dict[str, dict[str, str]] | None = None,
) -> list[str]:
    label_lines = []
    for roi_id in sorted(int(k) for k in labels if int(k) > 0):
        label = str(labels[roi_id]).replace(" ", "_")
        if label == "*.*.*.*.*":
            continue
        if hcp_region_info and label in hcp_region_info:
            label = hcp_region_info[label]["full_name"].replace(" ", "_")
        label_lines.append(f"{roi_id} {label} {roi_id}\n")
    return label_lines


def write_mricrogl_label_sidecars(
    out_dir: Path,
    stem: str,
    labels: dict[int, str],
    hcp_region_info: dict[str, dict[str, str]] | None = None,
) -> None:
    label_lines = mricrogl_label_lines(labels, hcp_region_info)
    for suffix in (".nii.txt", ".txt"):
        with (out_dir / f"{stem}{suffix}").open("w", encoding="utf-8", newline="\n") as f:
            f.writelines(label_lines)


def build_merged_groups(labels: dict[int, str]) -> dict[str, list[tuple[int, str]]]:
    groups: dict[str, list[tuple[int, str]]] = {}
    for roi_id, roi_name in labels.items():
        if int(roi_id) <= 0:
            continue
        name = str(roi_name)
        if name == "*.*.*.*.*":
            continue
        groups.setdefault(merged_roi_name(name), []).append((int(roi_id), name))
    return groups


def write_mricrogl_atlas_files(
    out_dir: Path,
    stem: str,
    atlas_img,
    labels: dict[int, str],
    hcp_region_info: dict[str, dict[str, str]] | None = None,
) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    data = np.rint(atlas_img.get_fdata()).astype(np.int16)
    header = atlas_img.header.copy()
    header.set_data_dtype(np.int16)
    header.set_intent("label")
    out_img = nib.Nifti1Image(data, atlas_img.affine, header)

    out_path = out_dir / f"{stem}.nii.gz"
    nib.save(out_img, str(out_path))

    # MRIcroGL commonly recognizes filename.nii.txt for filename.nii.gz atlases;
    # filename.txt is also written so the same label table is easy to find.
    write_mricrogl_label_sidecars(out_dir, stem, labels, hcp_region_info)


def save_mricrogl_atlas(
    atlas_name: str,
    atlas_img,
    labels: dict[int, str],
    hcp_region_info: dict[str, dict[str, str]] | None = None,
) -> None:
    short_name = atlas_short_name(atlas_name)
    if atlas_name.startswith("Schaefer"):
        alias_stems = ["Schaefer2018_400Parcels_7Networks_order_FSLMNI152_2mm"]
    elif atlas_name.startswith("HCP"):
        alias_stems = ["MNI_Glasser_HCP_v1.0"]
    elif atlas_name.startswith("AAL3"):
        alias_stems = ["AAL3v1"]
    elif atlas_name.startswith("Kastner") or atlas_name.startswith("Wang"):
        alias_stems = ["kastner"]
    elif atlas_name.startswith("HarvardOxford"):
        alias_stems = ["HarvardOxford_cort_sub_maxprob_thr25_2mm_split"]
    else:
        alias_stems = []
    atlas_out_dir = ATLAS_DIR / "MRIcroGL"
    label_info = hcp_region_info if atlas_name.startswith("HCP") else None
    for stem in [short_name, *alias_stems]:
        write_mricrogl_atlas_files(atlas_out_dir, stem, atlas_img, labels, label_info)

    if MICROGL_SYSTEM_ATLAS_DIR.exists():
        write_mricrogl_atlas_files(
            MICROGL_SYSTEM_ATLAS_DIR, short_name, atlas_img, labels, label_info
        )
        for stem in alias_stems:
            if (MICROGL_SYSTEM_ATLAS_DIR / f"{stem}.nii.gz").exists():
                write_mricrogl_label_sidecars(
                    MICROGL_SYSTEM_ATLAS_DIR, stem, labels, label_info
                )


def summarize_contrast_and_save_masks(
    base,
    contrast_name: str,
    contrast_path: Path,
    atlas_name: str,
    atlas_img,
    labels: dict[int, str],
    hcp_region_info: dict[str, dict[str, str]] | None = None,
) -> list[dict]:
    img = nib.load(str(contrast_path))
    data = img.get_fdata()
    atlas_data = np.rint(atlas_img.get_fdata()).astype(np.int32)
    voxel_volume = float(abs(np.linalg.det(img.affine[:3, :3])))
    groups = build_merged_groups(labels)

    out_mask_dir = OUT_DIR / atlas_short_name(atlas_name) / contrast_name
    out_mask_dir.mkdir(parents=True, exist_ok=True)
    for old_mask in out_mask_dir.glob("*.nii.gz"):
        old_mask.unlink()

    rows = []
    for merged_name in sorted(groups):
        members = groups[merged_name]
        member_ids = [roi_id for roi_id, _ in members]
        roi_mask = np.isin(atlas_data, member_ids)
        positive_roi_mask = roi_mask & (data > 0)
        positive_voxels = int(positive_roi_mask.sum())
        if positive_voxels == 0:
            continue

        merged_roi_total_voxels = int(roi_mask.sum())
        indices = np.column_stack(np.where(positive_roi_mask))
        weights = data[positive_roi_mask]
        center = base.weighted_center_mni(indices, weights, img.affine)

        masked_data = np.zeros(data.shape, dtype=np.float32)
        masked_data[positive_roi_mask] = data[positive_roi_mask].astype(np.float32)
        masked_img = nib.Nifti1Image(masked_data, img.affine, img.header)
        masked_img.set_data_dtype(np.float32)
        mask_path = out_mask_dir / f"{safe_filename(merged_name)}.nii.gz"
        nib.save(masked_img, str(mask_path))

        rows.append(
            {
                "contrast": contrast_name,
                "contrast_file": contrast_path.name,
                "atlas": atlas_name,
                "merged_roi_name": merged_name,
                "roi_full_name": readable_merged_roi_full_name(
                    merged_name, atlas_name, hcp_region_info
                ),
                "hcp_official_roi_name": (
                    hcp_region_info[merged_name]["official_name"]
                    if atlas_name.startswith("HCP")
                    and hcp_region_info
                    and merged_name in hcp_region_info
                    else ""
                ),
                "hcp_cortex_group": (
                    hcp_region_info[merged_name]["cortex"]
                    if atlas_name.startswith("HCP")
                    and hcp_region_info
                    and merged_name in hcp_region_info
                    else ""
                ),
                "hcp_lobe": (
                    hcp_region_info[merged_name]["lobe"]
                    if atlas_name.startswith("HCP")
                    and hcp_region_info
                    and merged_name in hcp_region_info
                    else ""
                ),
                "original_roi_names_merged": "; ".join(name for _, name in members),
                "roi_ids": "; ".join(str(roi_id) for roi_id in member_ids),
                "n_original_rois_merged": len(members),
                "positive_voxels_in_merged_roi": positive_voxels,
                "merged_roi_total_voxels": merged_roi_total_voxels,
                "percent_of_merged_roi_positive": (
                    100.0 * positive_voxels / merged_roi_total_voxels
                    if merged_roi_total_voxels
                    else np.nan
                ),
                "positive_volume_mm3": positive_voxels * voxel_volume,
                "merged_roi_volume_mm3": merged_roi_total_voxels * voxel_volume,
                "center_x_mni": center[0],
                "center_y_mni": center[1],
                "center_z_mni": center[2],
                "mean_positive_t": float(np.mean(weights)),
                "max_positive_t": float(np.max(weights)),
                "masked_stat_nii_gz": str(mask_path),
            }
        )
    return rows


def autosize_worksheet(ws) -> None:
    for column_cells in ws.columns:
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        if not values:
            continue
        width = min(max(max(len(value) for value in values) + 2, 10), 60)
        ws.column_dimensions[column_cells[0].column_letter].width = width


def write_workbook(summary_df: pd.DataFrame, all_df: pd.DataFrame, path: Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        for (contrast, atlas), group in all_df.groupby(["contrast", "atlas"]):
            contrast_short = (
                contrast.replace("AI_Pleasant_vs_AI_Neutral", "AI_Pleasant")
                .replace("Pleasant_vs_Neutral", "Pleasant")
                .replace("AI_Unpleasant_vs_AI_Neutral", "AI_Unpleasant")
                .replace("Unpleasant_vs_Neutral", "Unpleasant")
            )
            sheet_name = f"{contrast_short}_{atlas_short_name(atlas)}"[:31]
            group.to_excel(writer, sheet_name=sheet_name, index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            header_fill = PatternFill("solid", fgColor="D9EAF7")
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            autosize_worksheet(ws)


def main():
    base = load_base_module()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ATLAS_DIR.mkdir(parents=True, exist_ok=True)
    base.ATLAS_DIR = ATLAS_DIR

    missing = [str(path) for path in CONTRASTS.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing contrast map(s):\n" + "\n".join(missing))

    ref_img = nib.load(str(CONTRASTS["Pleasant_vs_Neutral"]))
    hcp_region_info = load_hcp_region_info()
    atlases = [
        base.load_schaefer(ref_img),
        base.load_hcp(ref_img),
        load_aal3(base, ref_img),
        load_kastner_wang(base, ref_img),
    ]
    for atlas_name, atlas_img, labels in atlases:
        save_mricrogl_atlas(atlas_name, atlas_img, labels, hcp_region_info)

    all_rows = []
    summary_rows = []
    for contrast_name, contrast_path in CONTRASTS.items():
        data = nib.load(str(contrast_path)).get_fdata()
        summary_rows.append(
            {
                "contrast": contrast_name,
                "contrast_file": contrast_path.name,
                "positive_voxels_total": int((data > 0).sum()),
                "negative_voxels_total": int((data < 0).sum()),
                "nonzero_voxels_total": int(np.count_nonzero(data)),
            }
        )
        for atlas_name, atlas_img, labels in atlases:
            rows = summarize_contrast_and_save_masks(
                base,
                contrast_name,
                contrast_path,
                atlas_name,
                atlas_img,
                labels,
                hcp_region_info,
            )
            all_rows.extend(rows)
            short_name = atlas_short_name(atlas_name)
            out_path = OUT_DIR / f"{contrast_name}_{short_name}_merged_positive_rois.csv"
            pd.DataFrame(rows).to_csv(out_path, index=False)
            print(f"Wrote {out_path} ({len(rows)} merged ROIs)")

    summary_df = pd.DataFrame(summary_rows)
    all_df = pd.DataFrame(all_rows)
    if not all_df.empty:
        all_df = all_df.sort_values(
            ["contrast", "atlas", "positive_voxels_in_merged_roi", "mean_positive_t"],
            ascending=[True, True, False, False],
        )

    all_path = OUT_DIR / "all_single_contrast_merged_positive_roi_tables.csv"
    summary_path = OUT_DIR / "single_contrast_merged_positive_roi_summary.csv"
    workbook_path = OUT_DIR / "roi_tables_28subs_fdr_positive_merged_rois.xlsx"
    all_df.to_csv(all_path, index=False)
    summary_df.to_csv(summary_path, index=False)
    write_workbook(summary_df, all_df, workbook_path)

    print(f"Wrote {all_path} ({len(all_df)} rows)")
    print(f"Wrote {summary_path}")
    print(f"Wrote {workbook_path}")
    print(f"Atlas files and labels saved under {ATLAS_DIR}")
    print(f"Masked positive-stat ROI images saved under {OUT_DIR / '<atlas>' / '<contrast>'}")


if __name__ == "__main__":
    main()
