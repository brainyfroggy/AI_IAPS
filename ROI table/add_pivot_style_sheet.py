from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROI_TABLE_DIR = Path(r"N:\Experimental_Data\yujunchen\projects\AI_IAPS\ROI table")
SOURCE_CSV = (
    ROI_TABLE_DIR
    / "all_single_contrast_merged_positive_roi_tables_HCP_Schaefer_AAL3_KastnerWang.csv"
)
WORKBOOKS = [
    ROI_TABLE_DIR / "roi_tables_28subs_fdr_positive_HCP_Schaefer_AAL3_KastnerWang.xlsx",
    ROI_TABLE_DIR
    / "28subs_fdr_positive_single_contrast_rois_merged_mricrogl"
    / "roi_tables_28subs_fdr_positive_merged_rois_readable_names.xlsx",
]
SHEET_NAME = "Pivot_style_summary"
CONTRASTS = [
    "Pleasant_vs_Neutral",
    "AI_Pleasant_vs_AI_Neutral",
    "Unpleasant_vs_Neutral",
    "AI_Unpleasant_vs_AI_Neutral",
]


def int_coord(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(int(float(value)))


def xyz_text(row: pd.Series) -> str:
    return (
        f"({int_coord(row['center_x_mni'])},"
        f"{int_coord(row['center_y_mni'])},"
        f"{int_coord(row['center_z_mni'])})"
    )


def build_wide_table(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["xyz_int"] = df.apply(xyz_text, axis=1)

    id_cols = ["atlas", "merged_roi_name", "roi_full_name"]
    keys = df[id_cols].drop_duplicates().sort_values(id_cols).reset_index(drop=True)
    wide = keys.copy()

    for contrast in CONTRASTS:
        sub = df[df["contrast"].eq(contrast)].copy()
        sub = sub[
            id_cols
            + [
                "contrast",
                "original_roi_names_merged",
                "xyz_int",
                "percent_of_merged_roi_positive",
            ]
        ]
        renamed = {
            "contrast": f"{contrast}_contrast",
            "original_roi_names_merged": f"{contrast}_original_roi_names_merged",
            "xyz_int": f"{contrast}_xyz_int",
            "percent_of_merged_roi_positive": (
                f"{contrast}_percent_of_merged_roi_positive"
            ),
        }
        sub = sub.rename(columns=renamed)
        wide = wide.merge(sub, how="left", on=id_cols)

    return wide


def autosize(ws) -> None:
    for column_cells in ws.columns:
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        if not values:
            continue
        width = min(max(max(len(value) for value in values) + 2, 10), 55)
        ws.column_dimensions[column_cells[0].column_letter].width = width


def write_sheet(workbook_path: Path, wide: pd.DataFrame) -> None:
    wb = load_workbook(workbook_path)
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME, 1)

    for col_idx, column in enumerate(wide.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=column)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row_idx, row in enumerate(wide.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            if pd.isna(value):
                value = ""
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if isinstance(value, float):
                cell.number_format = "0.0"

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)

    for idx, col_name in enumerate(wide.columns, start=1):
        if col_name.endswith("_percent_of_merged_roi_positive"):
            ws.column_dimensions[get_column_letter(idx)].width = 18
        elif col_name.endswith("_xyz_int"):
            ws.column_dimensions[get_column_letter(idx)].width = 16

    wb.save(workbook_path)


def main() -> None:
    df = pd.read_csv(SOURCE_CSV)
    wide = build_wide_table(df)
    for workbook_path in WORKBOOKS:
        write_sheet(workbook_path, wide)
        print(f"Updated {workbook_path} with {SHEET_NAME} ({len(wide)} rows)")


if __name__ == "__main__":
    main()
