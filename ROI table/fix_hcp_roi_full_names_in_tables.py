from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill


ROI_TABLE_DIR = Path(__file__).resolve().parent
OUT_DIRS = [
    ROI_TABLE_DIR / "28subs_fdr_positive_single_contrast_rois_merged_mricrogl",
    ROI_TABLE_DIR / "28subs_fdr_positive_single_contrast_rois_merged",
]
HCP_REGION_CSV = (
    ROI_TABLE_DIR
    / "28subs_fdr_positive_single_contrast_rois_merged_mricrogl"
    / "atlas"
    / "HCP-MMP1_UniqueRegionList.csv"
)


def atlas_short_name(atlas: object) -> str:
    atlas_name = str(atlas)
    if atlas_name.startswith("Schaefer"):
        return "Schaefer400"
    if atlas_name.startswith("HCP"):
        return "HCP_MMP1"
    if atlas_name.startswith("AAL3"):
        return "AAL3"
    if atlas_name.startswith("Kastner") or atlas_name.startswith("Wang"):
        return "KastnerWang"
    return re.sub(r"[^A-Za-z0-9._-]+", "_", atlas_name).strip("._")[:20] or "Atlas"


def clean_readable_name(value: object) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value).replace("_", " ")).strip()


def hcp_interpretable_name(
    official_name: str,
    region: str,
    hemi: str,
    cortex: str,
) -> str:
    side = {"L": "L", "R": "R"}.get(hemi, hemi)
    cortex_name = clean_readable_name(cortex)
    name_without_side = re.sub(rf"\s+{re.escape(side)}$", "", official_name).strip()
    name_without_area = re.sub(r"^Area\s+", "", name_without_side).strip()

    if not official_name.startswith("Area "):
        return official_name
    if cortex_name and name_without_area == region:
        return f"{cortex_name} - {region} {side}"
    if cortex_name and re.fullmatch(r"[A-Za-z0-9+\-.]+", name_without_area):
        return f"{cortex_name} - {name_without_area} {side}"
    return f"{name_without_area} {side}".strip()


def load_hcp_info() -> dict[str, dict[str, str]]:
    df = pd.read_csv(HCP_REGION_CSV)
    info: dict[str, dict[str, str]] = {}
    for _, row in df.iterrows():
        hemi = str(row.get("LR", "")).strip()
        region = str(row.get("region", "")).strip()
        if not hemi or not region:
            continue
        label = f"{hemi}_{region}"
        official = clean_readable_name(row.get("regionLongName", label))
        cortex = clean_readable_name(row.get("cortex", ""))
        info[label] = {
            "roi_full_name": hcp_interpretable_name(official, region, hemi, cortex),
            "hcp_official_roi_name": official,
            "hcp_cortex_group": cortex,
            "hcp_lobe": clean_readable_name(row.get("Lobe", "")),
        }
    return info


def ensure_after(df: pd.DataFrame, column: str, after: str) -> pd.DataFrame:
    columns = list(df.columns)
    if column not in columns or after not in columns:
        return df
    columns.remove(column)
    insert_at = columns.index(after) + 1
    columns.insert(insert_at, column)
    return df[columns]


def update_table(df: pd.DataFrame, hcp_info: dict[str, dict[str, str]]) -> pd.DataFrame:
    if "merged_roi_name" not in df.columns or "atlas" not in df.columns:
        return df

    for column in ["hcp_official_roi_name", "hcp_cortex_group", "hcp_lobe"]:
        if column not in df.columns:
            df[column] = ""

    is_hcp = df["atlas"].astype(str).str.startswith("HCP")
    for idx, merged_name in df.loc[is_hcp, "merged_roi_name"].astype(str).items():
        info = hcp_info.get(merged_name)
        if not info:
            continue
        df.at[idx, "roi_full_name"] = info["roi_full_name"]
        df.at[idx, "hcp_official_roi_name"] = info["hcp_official_roi_name"]
        df.at[idx, "hcp_cortex_group"] = info["hcp_cortex_group"]
        df.at[idx, "hcp_lobe"] = info["hcp_lobe"]

    for column in ["hcp_lobe", "hcp_cortex_group", "hcp_official_roi_name"]:
        df = ensure_after(df, column, "roi_full_name")
    return df


def autosize_worksheet(ws) -> None:
    for column_cells in ws.columns:
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        if not values:
            continue
        width = min(max(max(len(value) for value in values) + 2, 10), 60)
        ws.column_dimensions[column_cells[0].column_letter].width = width


def write_workbook(out_dir: Path, all_df: pd.DataFrame | None = None) -> None:
    all_path = out_dir / "all_single_contrast_merged_positive_roi_tables.csv"
    summary_path = out_dir / "single_contrast_merged_positive_roi_summary.csv"
    if (all_df is None and not all_path.exists()) or not summary_path.exists():
        return

    if all_df is None:
        all_df = pd.read_csv(all_path)
    summary_df = pd.read_csv(summary_path)
    workbook_paths = [
        out_dir / "roi_tables_28subs_fdr_positive_merged_rois.xlsx",
        out_dir / "all_single_contrast_merged_positive_roi_tables.xlsx",
        out_dir / "roi_tables_28subs_fdr_positive_merged_rois_readable_names.xlsx",
        out_dir / "all_single_contrast_merged_positive_roi_tables_readable_names.xlsx",
    ]

    for workbook_path in workbook_paths:
        try:
            writer_context = pd.ExcelWriter(workbook_path, engine="openpyxl")
        except PermissionError:
            print(f"Skipped locked workbook {workbook_path}")
            continue
        try:
            writer = writer_context.__enter__()
            summary_df.to_excel(writer, sheet_name="Summary", index=False)
            all_df.to_excel(writer, sheet_name="All_ROIs", index=False)
            for (contrast, atlas), group in all_df.groupby(["contrast", "atlas"]):
                contrast_short = (
                    str(contrast)
                    .replace("AI_Pleasant_vs_AI_Neutral", "AI_Pleasant")
                    .replace("Pleasant_vs_Neutral", "Pleasant")
                    .replace("AI_Unpleasant_vs_AI_Neutral", "AI_Unpleasant")
                    .replace("Unpleasant_vs_Neutral", "Unpleasant")
                )
                sheet_name = f"{contrast_short}_{atlas_short_name(atlas)}"[:31]
                group.to_excel(writer, sheet_name=sheet_name, index=False)

            for ws in writer.book.worksheets:
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                header_fill = PatternFill("solid", fgColor="D9EAF7")
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                    cell.fill = header_fill
                autosize_worksheet(ws)
        except PermissionError:
            print(f"Skipped locked workbook {workbook_path}")
        finally:
            try:
                writer_context.__exit__(None, None, None)
            except PermissionError:
                print(f"Skipped locked workbook {workbook_path}")


def main() -> None:
    hcp_info = load_hcp_info()
    for out_dir in OUT_DIRS:
        if not out_dir.exists():
            continue
        fixed_csv_dir = out_dir / "readable_name_fixed_csvs"
        csv_paths = sorted(out_dir.glob("*merged_positive_rois.csv"))
        csv_paths += [out_dir / "all_single_contrast_merged_positive_roi_tables.csv"]
        seen = set()
        fixed_all_df: pd.DataFrame | None = None
        for csv_path in csv_paths:
            if not csv_path.exists() or csv_path in seen:
                continue
            seen.add(csv_path)
            df = pd.read_csv(csv_path)
            df = update_table(df, hcp_info)
            try:
                df.to_csv(csv_path, index=False)
                print(f"Updated {csv_path}")
            except PermissionError:
                fixed_csv_dir.mkdir(parents=True, exist_ok=True)
                fixed_path = fixed_csv_dir / csv_path.name
                df.to_csv(fixed_path, index=False)
                print(f"Original CSV locked; wrote fixed copy {fixed_path}")
            if csv_path.name == "all_single_contrast_merged_positive_roi_tables.csv":
                fixed_all_df = df
        write_workbook(out_dir, fixed_all_df)
        print(f"Regenerated Excel workbooks in {out_dir}")


if __name__ == "__main__":
    main()
