from __future__ import annotations

import csv
import urllib.request
import xml.etree.ElementTree as ET
from pathlib import Path

import nibabel as nib
import numpy as np
import pandas as pd
from nilearn import datasets, image
from openpyxl.styles import Font, PatternFill


PROJECT_ROOT = Path(r"N:\Experimental_Data\yujunchen\projects\AI_IAPS")
CONTRAST_DIR = PROJECT_ROOT / "archive" / "outputs_seclvl"
OUT_DIR = PROJECT_ROOT / "ROI table"
ATLAS_DIR = OUT_DIR / "atlases"

CONTRASTS = {
    "Pleasant_vs_Neutral": CONTRAST_DIR / "Pl-Nt.nii",
    "PleasantAI_vs_NeutralAI": CONTRAST_DIR / "AI_Pl-AI_Nt.nii",
    "Unpleasant_vs_Neutral": CONTRAST_DIR / "Up-Nt.nii",
    "UnpleasantAI_vs_NeutralAI": CONTRAST_DIR / "AI_Up-AI_Nt.nii",
}

PAIRS = [
    (
        "Pleasant_vs_Neutral_AND_PleasantAI_vs_NeutralAI",
        "Pleasant_vs_Neutral",
        "PleasantAI_vs_NeutralAI",
    ),
    (
        "Unpleasant_vs_Neutral_AND_UnpleasantAI_vs_NeutralAI",
        "Unpleasant_vs_Neutral",
        "UnpleasantAI_vs_NeutralAI",
    ),
]

HCP_ATLAS_URL = (
    "https://raw.githubusercontent.com/mbedini/The-HCP-MMP1.0-atlas-in-FSL/"
    "master/MNI_Glasser_HCP_v1.0.nii.gz"
)
HCP_LABELS_URL = (
    "https://raw.githubusercontent.com/mbedini/The-HCP-MMP1.0-atlas-in-FSL/"
    "master/HCP-Multi-Modal-Parcellation-1.0.xml"
)


def download_if_missing(url: str, path: Path) -> None:
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    print(f"Downloading {path.name}...")
    urllib.request.urlretrieve(url, path)


def decode_label(label) -> str:
    if isinstance(label, bytes):
        return label.decode("utf-8")
    return str(label)


def load_schaefer(ref_img):
    fetched = datasets.fetch_atlas_schaefer_2018(
        n_rois=400,
        yeo_networks=7,
        resolution_mm=2,
        data_dir=str(ATLAS_DIR),
    )
    atlas_img = nib.load(fetched.maps)
    labels = {i: decode_label(label) for i, label in enumerate(fetched.labels)}
    labels.pop(0, None)
    resampled = image.resample_to_img(
        atlas_img,
        ref_img,
        interpolation="nearest",
        force_resample=True,
        copy_header=True,
    )
    return "Schaefer2018_400Parcels_7Networks", resampled, labels


def parse_hcp_labels(xml_path: Path) -> dict[int, str]:
    root = ET.parse(xml_path).getroot()
    labels: dict[int, str] = {}
    for elem in root.iter("label"):
        name = (elem.text or "").strip()
        if not name:
            continue
        idx_text = elem.attrib.get("index")
        if idx_text is None:
            continue
        idx = int(float(idx_text))
        labels[idx] = name
    return labels


def load_hcp(ref_img):
    atlas_path = ATLAS_DIR / "HCP-MMP1" / "MNI_Glasser_HCP_v1.0.nii.gz"
    labels_path = ATLAS_DIR / "HCP-MMP1" / "HCP-Multi-Modal-Parcellation-1.0.xml"
    download_if_missing(HCP_ATLAS_URL, atlas_path)
    download_if_missing(HCP_LABELS_URL, labels_path)

    atlas_img = nib.load(str(atlas_path))
    labels = parse_hcp_labels(labels_path)
    resampled = image.resample_to_img(
        atlas_img,
        ref_img,
        interpolation="nearest",
        force_resample=True,
        copy_header=True,
    )
    return "HCP_MMP1.0_Glasser_Volumetric", resampled, labels


def weighted_center_mni(indices: np.ndarray, weights: np.ndarray, affine: np.ndarray):
    coords = nib.affines.apply_affine(affine, indices)
    positive = np.asarray(weights, dtype=float)
    if not np.isfinite(positive).all() or positive.sum() <= 0:
        center = coords.mean(axis=0)
    else:
        center = np.average(coords, axis=0, weights=positive)
    return center


def summarize_pair(pair_name, map_a_name, map_b_name, atlas_name, atlas_img, labels):
    img_a = nib.load(str(CONTRASTS[map_a_name]))
    img_b = nib.load(str(CONTRASTS[map_b_name]))
    data_a = img_a.get_fdata()
    data_b = img_b.get_fdata()
    atlas_data = np.rint(atlas_img.get_fdata()).astype(np.int32)

    overlap = (data_a > 0) & (data_b > 0) & (atlas_data > 0)
    pair_weights = (data_a + data_b) / 2.0
    voxel_volume = float(abs(np.linalg.det(img_a.affine[:3, :3])))

    rows = []
    for roi_id in sorted(int(v) for v in np.unique(atlas_data[overlap]) if v > 0):
        roi_mask = overlap & (atlas_data == roi_id)
        n_vox = int(roi_mask.sum())
        if n_vox == 0:
            continue

        indices = np.column_stack(np.where(roi_mask))
        weights = pair_weights[roi_mask]
        center = weighted_center_mni(indices, weights, img_a.affine)

        atlas_roi_voxels = int((atlas_data == roi_id).sum())
        rows.append(
            {
                "comparison": pair_name,
                "atlas": atlas_name,
                "roi_id": roi_id,
                "roi_name": labels.get(roi_id, f"ROI_{roi_id}"),
                "overlap_voxels": n_vox,
                "overlap_volume_mm3": n_vox * voxel_volume,
                "percent_of_atlas_roi": (
                    100.0 * n_vox / atlas_roi_voxels if atlas_roi_voxels else np.nan
                ),
                "center_x_mni": center[0],
                "center_y_mni": center[1],
                "center_z_mni": center[2],
                "mean_pair_value": float(np.mean(weights)),
                "max_pair_value": float(np.max(weights)),
                f"mean_{map_a_name}_value": float(np.mean(data_a[roi_mask])),
                f"mean_{map_b_name}_value": float(np.mean(data_b[roi_mask])),
            }
        )
    return rows


def write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        return
    keys = []
    for row in rows:
        for key in row:
            if key not in keys:
                keys.append(key)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=keys, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def autosize_worksheet(ws) -> None:
    for column_cells in ws.columns:
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        if not values:
            continue
        width = min(max(max(len(value) for value in values) + 2, 10), 48)
        ws.column_dimensions[column_cells[0].column_letter].width = width


def write_excel_workbook(summary_df: pd.DataFrame, all_df: pd.DataFrame) -> Path:
    path = OUT_DIR / "roi_overlap_tables.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        if not all_df.empty:
            for (comparison, atlas), group in all_df.groupby(["comparison", "atlas"]):
                atlas_short = "Schaefer400" if atlas.startswith("Schaefer") else "HCP_MMP1"
                comparison_short = "Pleasant" if comparison.startswith("Pleasant") else "Unpleasant"
                sheet_name = f"{comparison_short}_{atlas_short}"[:31]
                group.to_excel(writer, sheet_name=sheet_name, index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            header_fill = PatternFill("solid", fgColor="D9EAF7")
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            autosize_worksheet(ws)
    return path


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ATLAS_DIR.mkdir(parents=True, exist_ok=True)

    ref_img = nib.load(str(CONTRASTS["Pleasant_vs_Neutral"]))
    atlases = [load_schaefer(ref_img), load_hcp(ref_img)]

    all_rows = []
    pair_summary_rows = []
    for pair_name, map_a_name, map_b_name in PAIRS:
        data_a = nib.load(str(CONTRASTS[map_a_name])).get_fdata()
        data_b = nib.load(str(CONTRASTS[map_b_name])).get_fdata()
        mask_a = data_a > 0
        mask_b = data_b > 0
        pair_summary_rows.append(
            {
                "comparison": pair_name,
                "map_a": map_a_name,
                "map_b": map_b_name,
                "map_a_nonzero_voxels": int(mask_a.sum()),
                "map_b_nonzero_voxels": int(mask_b.sum()),
                "overlap_nonzero_voxels": int((mask_a & mask_b).sum()),
                "dice_coefficient": float(
                    2 * (mask_a & mask_b).sum() / (mask_a.sum() + mask_b.sum())
                ),
            }
        )

        for atlas_name, atlas_img, labels in atlases:
            rows = summarize_pair(pair_name, map_a_name, map_b_name, atlas_name, atlas_img, labels)
            all_rows.extend(rows)
            short_name = "schaefer400" if atlas_name.startswith("Schaefer") else "hcp_mmp1"
            out_path = OUT_DIR / f"{pair_name}_{short_name}_roi_overlap.csv"
            write_csv(out_path, rows)
            print(f"Wrote {out_path} ({len(rows)} ROIs)")

    all_df = pd.DataFrame(all_rows)
    if not all_df.empty:
        all_df = all_df.sort_values(
            ["comparison", "atlas", "overlap_voxels", "mean_pair_value"],
            ascending=[True, True, False, False],
        )
        all_path = OUT_DIR / "all_roi_overlap_tables.csv"
        all_df.to_csv(all_path, index=False)
        print(f"Wrote {all_path} ({len(all_df)} rows)")

    summary_df = pd.DataFrame(pair_summary_rows)
    summary_path = OUT_DIR / "contrast_pair_overlap_summary.csv"
    summary_df.to_csv(summary_path, index=False)
    print(f"Wrote {summary_path}")

    workbook_path = write_excel_workbook(summary_df, all_df)
    print(f"Wrote {workbook_path}")


if __name__ == "__main__":
    main()
