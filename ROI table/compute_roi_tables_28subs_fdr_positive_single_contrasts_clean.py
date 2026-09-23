from __future__ import annotations

import importlib.util
from pathlib import Path

import nibabel as nib
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill


PROJECT_ROOT = Path(r"N:\Experimental_Data\yujunchen\projects\AI_IAPS")
ROI_TABLE_DIR = PROJECT_ROOT / "ROI table"
BASE_SCRIPT = ROI_TABLE_DIR / "compute_roi_overlap_tables.py"
CONTRAST_DIR = PROJECT_ROOT / "GLM" / "outputs_seclvl" / "28subs"
OUT_DIR = ROI_TABLE_DIR / "28subs_fdr_positive_single_contrast_rois_clean"

CONTRASTS = {
    "Pleasant_vs_Neutral": CONTRAST_DIR / "fdr_tstat_cond_Pl-Nt.nii",
    "AI_Pleasant_vs_AI_Neutral": CONTRAST_DIR / "fdr_tstat_cond_AI_Pl-AI_Nt.nii",
    "Unpleasant_vs_Neutral": CONTRAST_DIR / "fdr_tstat_cond_Up-Nt.nii",
    "AI_Unpleasant_vs_AI_Neutral": CONTRAST_DIR / "fdr_tstat_cond_AI_Up-AI_Nt.nii",
}


def load_base_module():
    spec = importlib.util.spec_from_file_location("roi_overlap_base", BASE_SCRIPT)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not load {BASE_SCRIPT}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def summarize_contrast(base, contrast_name, contrast_path, atlas_name, atlas_img, labels):
    img = nib.load(str(contrast_path))
    data = img.get_fdata()
    atlas_data = np.rint(atlas_img.get_fdata()).astype(np.int32)
    positive = (data > 0) & (atlas_data > 0)
    voxel_volume = float(abs(np.linalg.det(img.affine[:3, :3])))

    rows = []
    for roi_id in sorted(int(v) for v in np.unique(atlas_data[positive]) if v > 0):
        roi_mask = atlas_data == roi_id
        positive_roi_mask = positive & roi_mask
        positive_voxels = int(positive_roi_mask.sum())
        atlas_roi_total_voxels = int(roi_mask.sum())
        if positive_voxels == 0:
            continue

        indices = np.column_stack(np.where(positive_roi_mask))
        weights = data[positive_roi_mask]
        center = base.weighted_center_mni(indices, weights, img.affine)

        rows.append(
            {
                "contrast": contrast_name,
                "contrast_file": contrast_path.name,
                "atlas": atlas_name,
                "roi_id": roi_id,
                "roi_name": labels.get(roi_id, f"ROI_{roi_id}"),
                "positive_voxels_in_roi": positive_voxels,
                "atlas_roi_total_voxels": atlas_roi_total_voxels,
                "percent_of_atlas_roi_positive": (
                    100.0 * positive_voxels / atlas_roi_total_voxels
                    if atlas_roi_total_voxels
                    else np.nan
                ),
                "positive_volume_mm3": positive_voxels * voxel_volume,
                "atlas_roi_volume_mm3": atlas_roi_total_voxels * voxel_volume,
                "center_x_mni": center[0],
                "center_y_mni": center[1],
                "center_z_mni": center[2],
                "mean_positive_t": float(np.mean(weights)),
                "max_positive_t": float(np.max(weights)),
            }
        )
    return rows


def autosize_worksheet(ws) -> None:
    for column_cells in ws.columns:
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        if not values:
            continue
        width = min(max(max(len(value) for value in values) + 2, 10), 48)
        ws.column_dimensions[column_cells[0].column_letter].width = width


def write_workbook(summary_df: pd.DataFrame, all_df: pd.DataFrame, path: Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        for (contrast, atlas), group in all_df.groupby(["contrast", "atlas"]):
            atlas_short = "Schaefer400" if atlas.startswith("Schaefer") else "HCP_MMP1"
            contrast_short = (
                contrast.replace("AI_Pleasant_vs_AI_Neutral", "AI_Pleasant")
                .replace("Pleasant_vs_Neutral", "Pleasant")
                .replace("AI_Unpleasant_vs_AI_Neutral", "AI_Unpleasant")
                .replace("Unpleasant_vs_Neutral", "Unpleasant")
            )
            sheet_name = f"{contrast_short}_{atlas_short}"[:31]
            group.to_excel(writer, sheet_name=sheet_name, index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            header_fill = PatternFill("solid", fgColor="D9EAF7")
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            autosize_worksheet(ws)


def main():
    base = load_base_module()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    base.ATLAS_DIR = ROI_TABLE_DIR / "atlases"

    missing = [str(path) for path in CONTRASTS.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing contrast map(s):\n" + "\n".join(missing))

    ref_img = nib.load(str(CONTRASTS["Pleasant_vs_Neutral"]))
    atlases = [base.load_schaefer(ref_img), base.load_hcp(ref_img)]

    all_rows = []
    summary_rows = []
    for contrast_name, contrast_path in CONTRASTS.items():
        data = nib.load(str(contrast_path)).get_fdata()
        summary_rows.append(
            {
                "contrast": contrast_name,
                "contrast_file": contrast_path.name,
                "positive_voxels_total": int((data > 0).sum()),
                "negative_voxels_total": int((data < 0).sum()),
                "nonzero_voxels_total": int(np.count_nonzero(data)),
            }
        )
        for atlas_name, atlas_img, labels in atlases:
            rows = summarize_contrast(
                base, contrast_name, contrast_path, atlas_name, atlas_img, labels
            )
            all_rows.extend(rows)
            short_name = "schaefer400" if atlas_name.startswith("Schaefer") else "hcp_mmp1"
            out_path = OUT_DIR / f"{contrast_name}_{short_name}_positive_rois.csv"
            pd.DataFrame(rows).to_csv(out_path, index=False)
            print(f"Wrote {out_path} ({len(rows)} ROIs)")

    summary_df = pd.DataFrame(summary_rows)
    all_df = pd.DataFrame(all_rows)
    if not all_df.empty:
        all_df = all_df.sort_values(
            ["contrast", "atlas", "positive_voxels_in_roi", "mean_positive_t"],
            ascending=[True, True, False, False],
        )

    all_path = OUT_DIR / "all_single_contrast_positive_roi_tables_clean.csv"
    summary_path = OUT_DIR / "single_contrast_positive_roi_summary_clean.csv"
    workbook_path = OUT_DIR / "roi_tables_28subs_fdr_positive_single_contrasts_clean.xlsx"
    all_df.to_csv(all_path, index=False)
    summary_df.to_csv(summary_path, index=False)
    write_workbook(summary_df, all_df, workbook_path)

    print(f"Wrote {all_path} ({len(all_df)} rows)")
    print(f"Wrote {summary_path}")
    print(f"Wrote {workbook_path}")


if __name__ == "__main__":
    main()
